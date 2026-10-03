"""Dien mau Excel ISO goc va dung PDF tu cung noi dung."""
from copy import copy
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from decimal import Decimal
from xml.sax.saxutils import escape
import openpyxl
from openpyxl.cell.cell import MergedCell

ROOT = Path(__file__).resolve().parents[2] / 'form-mau'
RANKS = {'KHONG_CHON':'Không chọn','DU_PHONG':'Dự phòng','TIEU_CHUAN':'Tiêu chuẩn',
         'CHINH_YEU':'Chính yếu','CHIEN_LUOC':'Chiến lược','A':'A (cũ)','B':'B (cũ)','C':'C (cũ)'}
CRITERIA = ['diem_chat_luong','diem_giao_hang','diem_gia_ca','diem_tam_voc',
            'diem_thanh_toan','diem_dich_vu','diem_thoi_gian_hop_tac','diem_gia_tri_giao_dich']


def _value(cell,value):
    if isinstance(value,Decimal): value=float(value)
    if isinstance(value,(date,datetime)): value=value.strftime('%d/%m/%Y')
    cell.value=value
    if isinstance(value,str): cell.data_type='s'  # Du lieu nguoi dung khong phai cong thuc.


def _rows(sheet,start,rows):
    styles=[copy(c._style) for c in sheet[start]]
    for row in sheet.iter_rows(min_row=start):
        for cell in row:
            if not isinstance(cell,MergedCell):cell.value=None
    for i,values in enumerate(rows,start):
        for j,value in enumerate(values,1):
            cell=sheet.cell(i,j);cell._style=copy(styles[j-1]);_value(cell,value)
    end=max(start,start+len(rows)-1)
    if sheet.max_row>end:sheet.delete_rows(end+1,sheet.max_row-end)
    sheet.print_area=f'A1:{openpyxl.utils.get_column_letter(sheet.max_column)}{end}'


def workbook(form,data,year):
    number=int(form[2:]);path=next(ROOT.glob(f'{number}. *.xlsx'))
    book=openpyxl.load_workbook(path);s=book.active
    if form=='BM03':
        rows=[[i,r['ma_ncc'],r['ten'],r.get('phan_loai_ncc'),r.get('dia_chi'),
               r.get('nganh_nghe'),r.get('san_pham'),r.get('xuat_xu'),r.get('nguoi_lien_he'),
               r.get('sdt'),r.get('email'),r.get('ghi_chu')] for i,r in enumerate(data,1)]
        _rows(s,9,rows);s.print_title_rows='1:8'
    elif form=='BM06':
        r=data
        for cell,key in [('B7','id'),('B8','ten_ncc'),('B9','dia_chi'),('B10','nguoi_lien_he'),
                         ('B11','email'),('B12','mst'),('B13','nganh_nghe'),('B14','ten_hang')]:
            _value(s[cell],r.get(key))
        for i,key in enumerate(CRITERIA,18):
            score=r.get(key);_value(s[f'E{i}'],score)
            _value(s[f'F{i}'],float(score)*s[f'D{i}'].value if score is not None else None)
        _value(s['F26'],r.get('diem_tong'))
        s['B29']='';s['C29']='';s['F29']=''
        _value(s['B29'],'Kết luận: '+RANKS.get(r.get('xep_loai'),str(r.get('xep_loai') or 'Chưa xác định')))
        _value(s['B30'],r.get('ghi_chu'))
        _value(s['B31'],f"Điểm chuẩn hóa trên {r.get('trong_so_du_lieu') or 0}% hệ số có dữ liệu; tiêu chí trống không tính.")
        _value(s['B35'],r.get('ten_nguoi_cham') or r.get('nguoi_danh_gia'))
        _value(s['B36'],r.get('ten_nguoi_duyet') or r.get('nguoi_duyet'))
        _value(s['E35'],r.get('ngay_danh_gia'));_value(s['E36'],r.get('ngay_duyet'))
        for row in (35,36,37):s.row_dimensions[row].height=44
        s.print_area='A1:F37'
    elif form=='BM07':
        _value(s['C1'],f'TỔNG HỢP KẾT QUẢ ĐÁNH GIÁ NHÀ CUNG CẤP\nNăm {year}')
        groups={}
        for r in data:
            key=(r['id_ncc'],r['id_mat_hang_ncc'])
            row=groups.setdefault(key,[len(groups)+1,r['ten_ncc'],r.get('dia_chi'),r['ten_hang']]+[None]*12)
            day=r['ngay_danh_gia'];month=day.month if hasattr(day,'month') else date.fromisoformat(day).month
            row[month+3]=RANKS.get(r['xep_loai'],r['xep_loai'])
        _rows(s,7,list(groups.values()));s.print_title_rows='1:6'
    else:
        rows=[]
        for r in data:
            raw=(r.get('ket_qua') or '').upper()
            conclusion='Đạt' if raw in ('CHAP_NHAN','CHẤP NHẬN') else (
              'Không đạt' if raw in ('DOI_TRA','ĐỔI TRẢ','KHIEU_NAI','KHIẾU NẠI','GIAM_GIA','GIẢM GIÁ') else 'Chưa xác định')
            rows.append([r['ngay_nhan'],r.get('ten_ncc'),r['ten_hang'],r.get('ma_vat_tu'),
                         r.get('mo_ta'),r.get('huong_xu_ly'),f"{conclusion} ({r.get('ket_qua') or 'chưa ghi'})",r.get('nguoi_giam_sat')])
        _rows(s,7,rows);s.print_title_rows='1:6'
    s.page_setup.paperSize=s.PAPERSIZE_A4
    s.page_setup.orientation='portrait' if form=='BM06' else 'landscape'
    s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
    s.sheet_properties.pageSetUpPr.fitToPage=True
    s.oddFooter.center.text='Trang &P / &N'
    for row in s:
        for cell in row:
            if isinstance(cell.value,str) and cell.value.startswith('Số trang:'):
                _value(cell,'Số trang: theo chân trang')
    return book


def excel(book):
    output=BytesIO();book.save(output);return output.getvalue()


def pdf(book,form):
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,Image
    font='NccUnicode'
    if font not in pdfmetrics.getRegisteredFontNames():
        path=next((p for p in [Path('C:/Windows/Fonts/arial.ttf'),Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')] if p.exists()),None)
        if path is None:raise RuntimeError('Missing Unicode font for NCC PDF')
        pdfmetrics.registerFont(TTFont(font,str(path)))
    style=ParagraphStyle('NCC',fontName=font,fontSize=8 if form=='BM06' else 7,leading=10 if form=='BM06' else 9)
    title=ParagraphStyle('NCCTitle',parent=style,fontSize=12,leading=15)
    s=book.active;size=A4 if form=='BM06' else landscape(A4)
    width=size[0]-36
    columns=s.max_column;rows=s.max_row
    widths=[s.column_dimensions[openpyxl.utils.get_column_letter(i)].width or 13 for i in range(1,columns+1)]
    widths=[w/sum(widths)*width for w in widths]
    content=[]
    for row in s.iter_rows(max_row=rows,max_col=columns):
        content.append([Paragraph(escape(str(c.value if c.value is not None else '')).replace('\n','<br/>'),style) for c in row])
    title_col=1 if form=='BM06' else (3 if form=='BM03' else 2)
    content[0][title_col]=Paragraph(escape(str(s.cell(1,title_col+1).value or '')).replace('\n','<br/>'),title)
    for logo in s._images:
        anchor=logo.anchor._from
        max_width=sum(widths[:3]) if form=='BM03' else widths[0]+(widths[1] if form!='BM06' else 0)
        factor=min(max_width/logo.width,40/logo.height)
        content[anchor.row][anchor.col]=Image(BytesIO(logo._data()),width=logo.width*factor,height=logo.height*factor)
    heights=[44 if form=='BM06' and i in (34,35,36) else None for i in range(rows)]
    table=Table(content,colWidths=widths,rowHeights=heights,repeatRows=0 if form=='BM06' else (8 if form=='BM03' else 6))
    commands=[('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),0.3,colors.grey),
              ('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3)]
    for merge in s.merged_cells.ranges:
        commands.append(('SPAN',(merge.min_col-1,merge.min_row-1),(merge.max_col-1,merge.max_row-1)))
    table.setStyle(TableStyle(commands))
    output=BytesIO()
    def footer(canvas,doc):
        canvas.setFont(font,7);canvas.drawRightString(size[0]-18,10,f'{form} · Trang {doc.page}')
    SimpleDocTemplate(output,pagesize=size,leftMargin=18,rightMargin=18,topMargin=18,bottomMargin=24).build([table],onFirstPage=footer,onLaterPages=footer)
    return output.getvalue()

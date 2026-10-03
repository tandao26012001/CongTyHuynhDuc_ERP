"""API acceptance against configured DB; all fixture writes are rolled back.

Run from module/mua-hang with: python scripts/nghiem_thu_f1.py
Authentication sessions are simulated. Permissions, services and SQL are real.
Does not certify browser UAT, login, or purchasing transaction calculations.
"""
import sys
from pathlib import Path
from contextlib import ExitStack, contextmanager
from io import BytesIO
from uuid import uuid4
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from backend.api.app import app
from backend.api.middleware import auth_service
from backend.data import db


def main():
    passed = []
    fixture_ids = {}
    original = db.get_conn
    try:
        with original() as conn:
            with conn.transaction(force_rollback=True), ExitStack() as stack:
                @contextmanager
                def shared():
                    with conn.transaction():
                        yield conn

                for name, module in list(sys.modules.items()):
                    if name.startswith('backend.') and getattr(module, 'get_conn', None) is original:
                        stack.enter_context(patch.object(module, 'get_conn', shared))
                account = conn.execute('SELECT ma_tai_khoan,ma_nhan_vien,ma_bo_phan FROM tai_khoan LIMIT 1').fetchone()
                assert account, 'Need an existing account for actor foreign keys'
                profile = {**dict(account), 'vai_tro': 'ADMIN'}
                stack.enter_context(patch.object(auth_service, 'lay_ho_so', side_effect=lambda _: dict(profile)))
                client = TestClient(app)

                def request(label, method, path, expected=200, role='ADMIN', key=None, **kwargs):
                    profile['vai_tro'] = role
                    headers = {'X-Idempotency-Key': key or str(uuid4())}
                    response = client.request(method, '/api/v1' + path, headers=headers, **kwargs)
                    assert response.status_code == expected, f'{label}: HTTP {response.status_code}'
                    passed.append(label)
                    print('PASS', label, flush=True)
                    return response

                code = 'UAT' + uuid4().hex[:8]
                conn.execute('INSERT INTO chung_loai(ma_chung_loai,ten) VALUES(%s,%s)', (code, 'F1 synthetic acceptance'))
                fixture_ids['group'] = code
                catalog = request('classification catalog', 'GET', '/mat-hang-ncc/danh-muc').json()['data']
                assert catalog['don_vi_tinh']
                supplier = request('propose supplier', 'POST', '/nha-cung-cap/de-xuat', role='NV_YEU_CAU',
                                   json={'ten': 'UAT F1 ' + uuid4().hex, 'la_ncc_mua_hang': True,
                                         'xac_nhan_trung': True}).json()['data']['item']
                fixture_ids['supplier'] = supplier['id']
                assert supplier['trang_thai_xet_duyet'] == 'DE_XUAT'
                path = '/nha-cung-cap/' + supplier['id']
                request('supplier approval denied for requester', 'POST', path + '/duyet-de-xuat', 403,
                        role='NV_YEU_CAU', json={'phien_ban': supplier['phien_ban']})
                supplier = request('approve supplier', 'POST', path + '/duyet-de-xuat',
                                   role='NV_MUA_HANG', json={'phien_ban': supplier['phien_ban']}).json()['data']
                assert supplier['trang_thai_xet_duyet'] == 'DA_DUYET'
                assert supplier['da_phe_duyet'] is True, 'Approved supplier must be eligible for quotation requests'
                body = {'id_ncc': supplier['id'], 'ten_hang': 'F1 synthetic group', 'loai': 'HANG_HOA',
                        'nhom_hang_chinh': code, 'dvt': catalog['don_vi_tinh'][0]['dvt']}
                request('read-only cannot create group', 'POST', '/mat-hang-ncc', 403, role='CHI_XEM', json=body)
                key = str(uuid4())
                group = request('propose group', 'POST', '/mat-hang-ncc', role='NV_YEU_CAU', key=key, json=body).json()['data']
                fixture_ids['item'] = group['id']
                assert group['trang_thai'] == 'DE_XUAT' and group['pham_vi_danh_gia'] == 'NHOM_HANG'
                replay = request('group idempotency', 'POST', '/mat-hang-ncc', role='NV_YEU_CAU', key=key, json=body).json()['data']
                assert replay['id'] == group['id']
                request('duplicate group rejected', 'POST', '/mat-hang-ncc', 409, json=body)
                gpath = '/mat-hang-ncc/' + group['id']
                scores = dict.fromkeys(['diem_gia_ca', 'diem_tam_voc', 'diem_thanh_toan', 'diem_dich_vu'], 8)
                request('cannot score unapproved group', 'POST', gpath + '/danh-gia', 400, json=scores)
                edit = {**body, 'phien_ban': group['phien_ban'], 'thong_so_ky_thuat': 'Synthetic specification'}
                request('requester cannot edit group', 'PATCH', gpath, 403, role='NV_YEU_CAU', json=edit)
                group = request('edit group', 'PATCH', gpath, json=edit).json()['data']
                request('stale edit rejected', 'PATCH', gpath, 409, json=edit)
                group = request('approve group', 'POST', gpath + '/duyet', role='NV_MUA_HANG',
                                json={'phien_ban': group['phien_ban']}).json()['data']
                request('repeated group approval rejected', 'POST', gpath + '/duyet', 409,
                        json={'phien_ban': group['phien_ban']})
                history = request('group audit history', 'GET', gpath + '/lich-su').json()['data']
                assert len(history) == 3
                source = request('score source without transactions', 'GET', gpath + '/nguon-diem').json()['data']
                assert source['diem_xem_truoc']['diem_chat_luong'] is None
                key = str(uuid4())
                score = request('create assessment', 'POST', gpath + '/danh-gia', role='NV_MUA_HANG', key=key, json=scores).json()['data']
                assert float(score['diem_tong']) == 80 and score['xep_loai'] == 'CHINH_YEU'
                replay = request('assessment idempotency', 'POST', gpath + '/danh-gia', key=key, json=scores).json()['data']
                assert replay['id'] == score['id']
                request('duplicate pending assessment rejected', 'POST', gpath + '/danh-gia', 409, json=scores)
                request('employee cannot approve assessment', 'POST', '/danh-gia-ncc/' + score['id'] + '/duyet', 403,
                        role='NV_MUA_HANG', json={'phien_ban': score['phien_ban']})
                request('unapproved BM06 cannot export', 'GET', '/bieu-mau/ncc/BM06/tai', 400,
                        params={'id_danh_gia': score['id']})
                request('manager approves assessment', 'POST', '/danh-gia-ncc/' + score['id'] + '/duyet',
                        role='TBP_MUA_HANG', json={'phien_ban': score['phien_ban']})
                for form in ('BM03', 'BM06', 'BM07', 'BM08'):
                    for fmt in ('xlsx', 'pdf'):
                        params = {'id_ncc': supplier['id'], 'dinh_dang': fmt}
                        if form == 'BM06':
                            params['id_danh_gia'] = score['id']
                        response = request(form + ' ' + fmt, 'GET', '/bieu-mau/ncc/' + form + '/tai', params=params)
                        if fmt == 'xlsx':
                            wb = load_workbook(BytesIO(response.content))
                            assert wb.active.max_row > 0
                        else:
                            assert response.content.startswith(b'%PDF-') and b'%%EOF' in response.content[-100:]
                assert conn.execute('SELECT count(*) AS n FROM ncc_bieu_mau_xuat WHERE bo_loc->>\'id_ncc\'=%s',
                                    (supplier['id'],)).fetchone()['n'] == 8
                request('BM08 fixed route', 'GET', '/nha-cung-cap/so-bm08')
                request('due assessment fixed route', 'GET', '/nha-cung-cap/danh-gia-den-han')
                limit = request('monthly limit available', 'GET', path + '/dinh-muc').json()['data']
                limit_body = {'phien_ban': limit['phien_ban'], 'dinh_muc_thang': 1000000}
                request('employee cannot set monthly limit', 'PATCH', path + '/dinh-muc', 403,
                        role='NV_MUA_HANG', json=limit_body)
                limit = request('manager sets monthly limit', 'PATCH', path + '/dinh-muc',
                                role='TBP_MUA_HANG', json=limit_body).json()['data']
                visible = request('admin reads saved monthly limit', 'GET', path + '/dinh-muc').json()['data']
                assert float(visible['con_lai']) == 1000000
                request('stale monthly limit update rejected', 'PATCH', path + '/dinh-muc', 409, json=limit_body)
                hidden = request('requester cannot see monetary limit', 'GET', path + '/dinh-muc',
                                 role='NV_YEU_CAU').json()['data']
                assert hidden['dinh_muc_thang'] is None and hidden['da_dat_thang_nay'] is None
                request('score outside allowed range rejected', 'POST', gpath + '/danh-gia', 400,
                        json={**scores, 'diem_gia_ca': 11})
                request('new assessment after approval', 'POST', gpath + '/danh-gia', json=scores)
            for table, column, value in [('chung_loai', 'ma_chung_loai', fixture_ids['group']),
                                          ('nha_cung_cap', 'id', fixture_ids['supplier']),
                                          ('mat_hang_ncc', 'id', fixture_ids['item'])]:
                assert conn.execute(f'SELECT count(*) AS n FROM {table} WHERE {column}=%s', (value,)).fetchone()['n'] == 0
            print(f'PASS rollback verified; {len(passed)} API checks')
    finally:
        db.close_pool()


if __name__ == '__main__':
    main()
